"""Offline (no-AWS) service-level smoke test for the Phase 2 pipeline.

Run INSIDE the api/worker container (it has tesseract, Tika/Java, spaCy, etc.):

    docker compose exec -T worker python -m tests.test_offline

It generates fixtures (text-layer PDF via reportlab, typed PNG via Pillow, DOCX via
python-docx) into tests/fixtures/ (host-visible through the ./backend bind mount), then
exercises the non-AWS code paths: parser/OCR extraction, rule-based classification, Fernet
field encryption round-trip, and local storage with path-traversal protection.
"""
import sys
import traceback
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURES.mkdir(parents=True, exist_ok=True)

results: list[tuple[str, bool, str]] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    results.append((name, bool(cond), detail))
    mark = "PASS" if cond else "FAIL"
    print(f"[{mark}] {name}" + (f" — {detail}" if detail else ""))


# ── Fixture generation ────────────────────────────────────────────
def make_docx() -> Path:
    import docx
    path = FIXTURES / "sample_contract.docx"
    d = docx.Document()
    d.add_heading("Master Services Agreement", level=1)
    d.add_paragraph(
        "WHEREAS the parties agree to the terms herein, and subject to the governing law of "
        "the State of Delaware, each party shall indemnify the other against any breach. This "
        "agreement includes a termination clause effective upon thirty days written notice."
    )
    d.save(path)
    return path


def make_phi_docx() -> Path:
    """A document containing PHI (SSN + DOB) used by the Phase 3 HTTP QA upload."""
    import docx
    path = FIXTURES / "sample_phi.docx"
    d = docx.Document()
    d.add_heading("Patient Intake Record", level=1)
    d.add_paragraph(
        "Patient: Jane Roe. SSN 123-45-6789. Date of birth 01/02/1990. "
        "This record is shared with the care team. No consent form is currently on file."
    )
    d.save(path)
    return path


def make_png() -> Path:
    from PIL import Image, ImageDraw, ImageFont
    path = FIXTURES / "sample_invoice.png"
    img = Image.new("RGB", (1400, 360), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default(size=40)
    except TypeError:  # very old Pillow
        font = ImageFont.load_default()
    lines = [
        "INVOICE NUMBER 100245",
        "Amount Due: 4,200.00 USD",
        "Bill To: Acme Corp",
        "Payment Terms: Net 30",
    ]
    for i, line in enumerate(lines):
        draw.text((40, 30 + i * 75), line, fill="black", font=font)
    img.save(path)
    return path


def make_pdf() -> Path:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    path = FIXTURES / "sample_policy.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    text = c.beginText(72, 720)
    for line in [
        "PRIVACY POLICY",
        "This policy describes how we handle personal data and the rights of each data subject.",
        "Data controller responsibilities include honoring every right to erasure request.",
        "We retain personal data only as long as necessary for the stated purpose.",
    ]:
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()
    return path


# ── Tests ─────────────────────────────────────────────────────────
def run() -> None:
    from app.services.ocr_service import extract_text
    from app.services import classifier_service
    from app.utils.encryption import encrypt_text, decrypt_text
    from app.services import storage_service

    # DOCX (python-docx) — fully offline, deterministic
    docx_path = make_docx()
    docx_text = extract_text(str(docx_path), "docx")
    check("DOCX extraction returns text", len(docx_text) > 50, f"{len(docx_text)} chars")
    docx_cls = classifier_service.classify(docx_text)
    check("DOCX classified as Legal Contract",
          docx_cls.get("document_type") == "Legal Contract",
          f"{docx_cls.get('document_type')} @ {docx_cls.get('confidence')}")

    # PHI fixture (for the Phase 3 HTTP QA) — generate + sanity-check it contains the SSN
    phi_path = make_phi_docx()
    phi_text = extract_text(str(phi_path), "docx")
    check("PHI DOCX contains SSN", "123-45-6789" in phi_text, f"{len(phi_text)} chars")

    # PNG (Tesseract) — offline OCR
    png_path = make_png()
    png_text = extract_text(str(png_path), "png")
    check("PNG OCR returns text", len(png_text.strip()) > 0, repr(png_text[:80]))
    check("PNG OCR picked up 'INVOICE'", "invoice" in png_text.lower(), repr(png_text[:80]))

    # PDF (Tika) — needs Java (in image); no AWS
    pdf_path = make_pdf()
    pdf_text = extract_text(str(pdf_path), "pdf")
    check("PDF (Tika) extraction returns text", len(pdf_text.strip()) > 0, f"{len(pdf_text)} chars")

    # Fernet field encryption round-trip
    secret = "Patient SSN 123-45-6789 — confidential."
    ct = encrypt_text(secret)
    check("extracted_text encrypts to Fernet ciphertext", ct.startswith("gAAA") and ct != secret)
    check("Fernet decrypt round-trips", decrypt_text(ct) == secret)

    # Local storage save/read/delete + path-traversal rejection
    rel = "documents/_selftest/hello.txt"
    storage_service.save_file(b"hello world", rel)
    check("storage save+read round-trip", storage_service.read_file(rel) == b"hello world")
    traversal_blocked = False
    try:
        storage_service.get_file_path("../../etc/passwd")
    except ValueError:
        traversal_blocked = True
    check("path traversal (../) rejected", traversal_blocked)
    storage_service.delete_file(rel)


if __name__ == "__main__":
    try:
        run()
    except Exception:
        traceback.print_exc()
        results.append(("test harness crashed", False, ""))

    passed = sum(1 for _, ok, _ in results if ok)
    total = len(results)
    print(f"\n==== Offline QA: {passed}/{total} passed ====")
    sys.exit(0 if passed == total else 1)
