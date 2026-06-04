"""Audit report generation — PDF (ReportLab) and DOCX (python-docx).

Both formats are driven from the same document + checklist data and produce the 8 sections
defined in the spec. Returns raw file bytes; the router persists them to local storage and
serves them through an authenticated download endpoint.
"""
import io
from datetime import datetime, timezone

# ── shared helpers ────────────────────────────────────────────────
_SEVERITY_ORDER = {"critical": 0, "warning": 1, "info": 2}
_SEVERITY_COLOR = {"critical": "#991B1B", "warning": "#92400E", "info": "#1E3A8A"}


def _report_id(document: dict) -> str:
    return f"DG-{str(document.get('_id', ''))[-8:].upper() or 'PREVIEW'}"


def _sorted_violations(document: dict) -> list[dict]:
    return sorted(
        document.get("violations", []) or [],
        key=lambda v: _SEVERITY_ORDER.get(v.get("severity"), 3),
    )


def _severity_counts(violations: list[dict]) -> dict:
    counts = {"critical": 0, "warning": 0, "info": 0}
    for v in violations:
        counts[v.get("severity", "info")] = counts.get(v.get("severity", "info"), 0) + 1
    return counts


def _doc_info_rows(document: dict) -> list[tuple[str, str]]:
    conf = document.get("confidence_score")
    return [
        ("Original filename", str(document.get("original_name", "—"))),
        ("File type", str(document.get("file_type", "—")).upper()),
        ("Classification", str(document.get("classification", "—"))),
        ("Confidence", f"{round(conf * 100)}%" if isinstance(conf, (int, float)) else "—"),
        ("Status", str(document.get("status", "—"))),
        ("Approval", str(document.get("approval_status", "pending"))),
    ]


# ── PDF (ReportLab) ───────────────────────────────────────────────
def _build_pdf(document: dict, checklist: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=0.7 * inch,
                            bottomMargin=0.7 * inch, leftMargin=0.7 * inch, rightMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#1B2A4A")
    h = ParagraphStyle("H", parent=styles["Heading2"], textColor=navy, spaceBefore=14, spaceAfter=6)
    body = styles["BodyText"]
    small = ParagraphStyle("S", parent=body, fontSize=8, textColor=colors.HexColor("#64748B"))
    flow = []

    violations = _sorted_violations(document)
    counts = _severity_counts(violations)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    # 1. Header
    flow.append(Paragraph("DocuGuard — Compliance Audit Report", ParagraphStyle(
        "T", parent=styles["Title"], textColor=navy)))
    flow.append(Paragraph(f"Report ID: {_report_id(document)} &nbsp;&nbsp;|&nbsp;&nbsp; Generated: {now}", small))
    flow.append(Spacer(1, 10))

    # 2. Document Information
    flow.append(Paragraph("1. Document Information", h))
    info = Table([[k, v] for k, v in _doc_info_rows(document)], colWidths=[2 * inch, 4.5 * inch])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
        ("TEXTCOLOR", (0, 0), (0, -1), navy),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    flow.append(info)

    # 3. AI Summary
    flow.append(Paragraph("2. AI Summary", h))
    flow.append(Paragraph(document.get("summary") or "No summary available.", body))
    clauses = document.get("key_clauses") or []
    if clauses:
        flow.append(Paragraph("Key clauses:", ParagraphStyle("KC", parent=body, fontName="Helvetica-Bold")))
        for c in clauses:
            flow.append(Paragraph(f"• {c}", body))
    flow.append(Paragraph(f"<b>Purpose:</b> {document.get('document_purpose') or '—'}", body))

    # 4. Compliance Framework
    flow.append(Paragraph("3. Compliance Framework", h))
    flow.append(Paragraph(
        f"Checklist: <b>{checklist.get('name', document.get('checklist_id', '—'))}</b> "
        f"(version {checklist.get('version', '—')}) — {len(checklist.get('rules', []))} rules checked.", body))

    # 5. Violations Summary + risk score
    flow.append(Paragraph("4. Violations Summary", h))
    score = document.get("risk_score") or 0
    level = document.get("risk_level") or "—"
    bar = "█" * (score // 5) + "░" * (20 - score // 5)
    flow.append(Paragraph(
        f"Critical: <b>{counts['critical']}</b> &nbsp; Warning: <b>{counts['warning']}</b> &nbsp; "
        f"Info: <b>{counts['info']}</b>", body))
    flow.append(Paragraph(f"Risk score: <b>{score}/100</b> ({level})  [{bar}]", body))

    # 6. Detailed Violations Table
    flow.append(Paragraph("5. Detailed Violations", h))
    if violations:
        rows = [["Rule", "Category", "Severity", "Description", "AI Fix Suggestion"]]
        for v in violations:
            rows.append([
                v.get("rule_id", ""), v.get("category", ""), (v.get("severity") or "").upper(),
                Paragraph(v.get("description") or "", small),
                Paragraph(v.get("ai_fix_suggestion") or v.get("generic_fix") or "", small),
            ])
        vt = Table(rows, colWidths=[0.8 * inch, 1.0 * inch, 0.7 * inch, 1.9 * inch, 2.1 * inch], repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i, v in enumerate(violations, start=1):
            style.append(("TEXTCOLOR", (2, i), (2, i), colors.HexColor(_SEVERITY_COLOR.get(v.get("severity"), "#1E293B"))))
        vt.setStyle(TableStyle(style))
        flow.append(vt)
    else:
        flow.append(Paragraph("No violations detected.", body))

    # 7. Approval Status
    flow.append(Paragraph("6. Approval Status", h))
    flow.append(Paragraph(
        f"Status: <b>{document.get('approval_status', 'pending')}</b><br/>"
        f"Decision by: {document.get('approved_by') or '—'}<br/>"
        f"Reason: {document.get('rejection_reason') or '—'}", body))

    # 8. Auditor signature line
    flow.append(Spacer(1, 28))
    flow.append(Paragraph("7. Auditor Sign-off", h))
    flow.append(Spacer(1, 22))
    sig = Table([["Auditor signature: ______________________________", "Date: ______________"]],
                colWidths=[4.2 * inch, 2.3 * inch])
    sig.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("TEXTCOLOR", (0, 0), (-1, -1), navy)]))
    flow.append(sig)

    doc.build(flow)
    return buf.getvalue()


# ── DOCX (python-docx) ────────────────────────────────────────────
def _build_docx(document: dict, checklist: dict) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = docx.Document()
    violations = _sorted_violations(document)
    counts = _severity_counts(violations)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    navy = RGBColor(0x1B, 0x2A, 0x4A)

    title = d.add_heading("DocuGuard — Compliance Audit Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = d.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Report ID: {_report_id(document)}  |  Generated: {now}")
    run.font.size = Pt(9)

    d.add_heading("1. Document Information", level=1)
    info = d.add_table(rows=0, cols=2)
    info.style = "Light Grid Accent 1"
    for k, v in _doc_info_rows(document):
        cells = info.add_row().cells
        cells[0].text, cells[1].text = k, v

    d.add_heading("2. AI Summary", level=1)
    d.add_paragraph(document.get("summary") or "No summary available.")
    for c in (document.get("key_clauses") or []):
        d.add_paragraph(c, style="List Bullet")
    p = d.add_paragraph()
    p.add_run("Purpose: ").bold = True
    p.add_run(document.get("document_purpose") or "—")

    d.add_heading("3. Compliance Framework", level=1)
    d.add_paragraph(
        f"Checklist: {checklist.get('name', document.get('checklist_id', '—'))} "
        f"(version {checklist.get('version', '—')}) — {len(checklist.get('rules', []))} rules checked.")

    d.add_heading("4. Violations Summary", level=1)
    score = document.get("risk_score") or 0
    d.add_paragraph(
        f"Critical: {counts['critical']}   Warning: {counts['warning']}   Info: {counts['info']}")
    bar = "█" * (score // 5) + "░" * (20 - score // 5)
    d.add_paragraph(f"Risk score: {score}/100 ({document.get('risk_level') or '—'})  [{bar}]")

    d.add_heading("5. Detailed Violations", level=1)
    if violations:
        table = d.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, label in enumerate(["Rule", "Category", "Severity", "Description", "AI Fix Suggestion"]):
            hdr[i].text = label
            hdr[i].paragraphs[0].runs[0].bold = True
        for v in violations:
            cells = table.add_row().cells
            cells[0].text = v.get("rule_id", "")
            cells[1].text = v.get("category", "") or ""
            cells[2].text = (v.get("severity") or "").upper()
            cells[3].text = v.get("description") or ""
            cells[4].text = v.get("ai_fix_suggestion") or v.get("generic_fix") or ""
    else:
        d.add_paragraph("No violations detected.")

    d.add_heading("6. Approval Status", level=1)
    d.add_paragraph(f"Status: {document.get('approval_status', 'pending')}")
    d.add_paragraph(f"Decision by: {document.get('approved_by') or '—'}")
    d.add_paragraph(f"Reason: {document.get('rejection_reason') or '—'}")

    d.add_heading("7. Auditor Sign-off", level=1)
    d.add_paragraph("\nAuditor signature: ______________________________      Date: ______________")

    for para in d.paragraphs:
        if para.style.name.startswith("Heading"):
            for r in para.runs:
                r.font.color.rgb = navy

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def generate_report(document: dict, checklist: dict, fmt: str = "pdf") -> bytes:
    if fmt == "docx":
        return _build_docx(document, checklist or {})
    return _build_pdf(document, checklist or {})
