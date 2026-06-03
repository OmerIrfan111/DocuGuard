"""Text extraction for native (non-OCR) formats: Apache Tika (PDF) and python-docx (DOCX)."""


def tika_extract(file_path: str) -> str:
    """Extract text from a PDF (or other rich doc) via Apache Tika.

    Tika requires Java (installed in the backend image) and fetches its server jar on first
    use. Returns an empty string on failure so callers can fall back to Textract.
    """
    try:
        from tika import parser
        parsed = parser.from_file(str(file_path))
        return (parsed.get("content") or "").strip()
    except Exception:
        return ""


def docx_extract(file_path: str) -> str:
    """Extract text from a .docx using python-docx (paragraphs + table cells)."""
    try:
        import docx
        document = docx.Document(str(file_path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts).strip()
    except Exception:
        return ""
